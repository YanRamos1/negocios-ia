"""Testes da planilha de precificação de confeitaria.

Rodar com:
    python -m unittest discover -s produtos/planilhas/precificacao-confeitaria -p "test_*.py"
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import gerar_planilha_confeitaria as gerador  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

try:  # pragma: no cover - dependência opcional
    import formulas
except ImportError:  # pragma: no cover
    formulas = None

ABAS_COMPLETAS = {"Leia-me", "Config", "Ingredientes", "Itens", "Receitas", "Resumo"}
ABAS_GRATIS = {"Leia-me", "Config", "Ingredientes"}


def escalar(valor: object) -> object:
    if hasattr(valor, "tolist"):
        valor = valor.tolist()
    while isinstance(valor, list):
        if not valor:
            return 0
        valor = valor[0]
    return valor


def avaliar(caminho: Path) -> dict[str, object]:
    modelo = formulas.ExcelModel().loads(str(caminho)).finish()
    solucao = modelo.calculate()
    valores: dict[str, object] = {}
    for chave, resultado in solucao.items():
        if "!" not in chave:
            continue
        aba, celula = chave.rsplit("!", 1)
        nome = aba.rsplit("]", 1)[-1].strip("'").upper()
        endereco = celula.strip("'")
        valores[f"{nome}!{endereco}"] = escalar(getattr(resultado, "value", resultado))
    return valores


class EstruturaTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_versao_completa_tem_todas_as_abas(self) -> None:
        destino = gerador.gerar(self.dir / "completa.xlsx", completa=True)
        self.assertEqual(set(load_workbook(destino).sheetnames), ABAS_COMPLETAS)

    def test_versao_gratis_nao_tem_preco(self) -> None:
        destino = gerador.gerar(self.dir / "gratis.xlsx", completa=False)
        self.assertEqual(set(load_workbook(destino).sheetnames), ABAS_GRATIS)

    def test_ingredientes_tem_validacao_de_unidade(self) -> None:
        destino = gerador.gerar(self.dir / "completa.xlsx", completa=True)
        ws = load_workbook(destino)["Ingredientes"]
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 1)

    def test_receitas_deixam_linhas_em_branco(self) -> None:
        destino = gerador.gerar(self.dir / "completa.xlsx", completa=True)
        ws = load_workbook(destino)["Receitas"]
        self.assertEqual(ws["A4"].value, "Brigadeiro gourmet")
        self.assertIsNone(ws["A7"].value)


@unittest.skipIf(formulas is None, "biblioteca formulas não instalada")
class CalculoTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        destino = gerador.gerar(self.dir / "completa.xlsx", completa=True)
        self.valores = avaliar(destino)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_custo_unitario_do_ingrediente(self) -> None:
        self.assertAlmostEqual(self.valores["INGREDIENTES!E4"], 5.0, places=2)
        self.assertAlmostEqual(self.valores["INGREDIENTES!E7"], 1.0, places=2)

    def test_custo_dos_ingredientes_da_receita(self) -> None:
        self.assertAlmostEqual(self.valores["RECEITAS!D4"], 21.25, places=2)
        self.assertAlmostEqual(self.valores["RECEITAS!D5"], 10.71, places=2)

    def test_mao_de_obra_e_rateio_por_hora(self) -> None:
        self.assertAlmostEqual(self.valores["RECEITAS!E4"], 45.0, places=2)
        self.assertAlmostEqual(self.valores["RECEITAS!G4"], 11.025, places=3)

    def test_custo_total_e_porcao(self) -> None:
        self.assertAlmostEqual(self.valores["RECEITAS!H4"], 92.275, places=2)
        self.assertAlmostEqual(self.valores["RECEITAS!I4"], 3.0758, places=3)

    def test_preco_aplica_margem_e_taxa(self) -> None:
        self.assertAlmostEqual(self.valores["RECEITAS!J4"], 5.1264, places=3)
        self.assertAlmostEqual(self.valores["RECEITAS!K4"], 1.5379, places=3)

    def test_resumo_conta_receitas_e_margem(self) -> None:
        self.assertAlmostEqual(self.valores["RESUMO!B4"], 3, places=0)
        self.assertAlmostEqual(self.valores["RESUMO!B7"], 0.30, places=3)


if __name__ == "__main__":
    unittest.main()

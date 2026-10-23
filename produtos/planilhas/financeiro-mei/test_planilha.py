"""Testes da planilha Financeiro PF x PJ.

Rodar com:
    python -m unittest discover -s produtos/planilhas/financeiro-mei -p "test_*.py"

A validação de fórmulas usa a biblioteca `formulas` (opcional). Sem ela, os
testes de estrutura continuam rodando.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import gerar_planilha  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

try:  # pragma: no cover - dependência opcional
    import formulas
except ImportError:  # pragma: no cover
    formulas = None

ABAS_COMPLETAS = {"Leia-me", "Config", "Lançamentos", "Resumo", "Tabelas", "Comparativo"}
ABAS_GRATIS = {"Leia-me", "Config", "Lançamentos", "Resumo"}


def avaliar(caminho: Path) -> dict[str, float | str]:
    """Calcula o workbook e devolve as células da aba Comparativo por endereço."""
    modelo = formulas.ExcelModel().loads(str(caminho)).finish()
    solucao = modelo.calculate()
    valores: dict[str, float | str] = {}
    for chave, resultado in solucao.items():
        if "COMPARATIVO" not in chave.upper():
            continue
        endereco = chave.rsplit("!", 1)[-1].strip("'")
        conteudo = getattr(resultado, "value", resultado)
        if isinstance(conteudo, list) and conteudo and isinstance(conteudo[0], list):
            conteudo = conteudo[0][0]
        valores[endereco] = conteudo
    return valores


class EstruturaTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_versao_completa_tem_todas_as_abas(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        wb = load_workbook(destino)
        self.assertEqual(set(wb.sheetnames), ABAS_COMPLETAS)

    def test_versao_gratis_nao_tem_calculo(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "gratis.xlsx", completa=False)
        wb = load_workbook(destino)
        self.assertEqual(set(wb.sheetnames), ABAS_GRATIS)

    def test_tabelas_comecam_em_zero(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        ws = load_workbook(destino)["Tabelas"]
        self.assertEqual(ws["A5"].value, 0)
        self.assertEqual(ws["A11"].value, 0)
        self.assertEqual(ws["A18"].value, 0)
        self.assertEqual(ws["A26"].value, 0)

    def test_lancamentos_tem_validacao_de_lista(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        ws = load_workbook(destino)["Lançamentos"]
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 4)


@unittest.skipIf(formulas is None, "biblioteca formulas não instalada")
class CalculoTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_cenario_padrao_bate_com_o_site(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        valores = avaliar(destino)
        self.assertAlmostEqual(valores["B6"], 501.51, places=2)   # INSS CLT
        self.assertAlmostEqual(valores["B7"], 0, places=2)        # IRRF CLT
        self.assertAlmostEqual(valores["B10"], 4498.49, places=2) # líquido CLT
        self.assertAlmostEqual(valores["C5"], 1820.00, places=2)  # pró-labore
        self.assertAlmostEqual(valores["C6"], 200.20, places=2)   # INSS pró-labore
        self.assertAlmostEqual(valores["C8"], 390.00, places=2)   # DAS 6%
        self.assertAlmostEqual(valores["C10"], 5609.80, places=2) # líquido PJ
        self.assertEqual(valores["B13"], "PJ rende mais")

    def test_config_alterada_recalcula(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        wb = load_workbook(destino)
        wb["Config"]["B4"] = 3000.00
        wb["Config"]["B5"] = 4000.00
        wb.save(destino)

        valores = avaliar(destino)
        self.assertAlmostEqual(valores["B10"], 2751.40, places=2)
        self.assertAlmostEqual(valores["C10"], 3281.69, places=2)

    def test_anexo_v_quando_fator_r_abaixo_de_28(self) -> None:
        destino = gerar_planilha.gerar(self.dir / "completa.xlsx", completa=True)
        wb = load_workbook(destino)
        wb["Config"]["B6"] = 0.10
        wb.save(destino)

        valores = avaliar(destino)
        self.assertAlmostEqual(valores["F16"], 0.10, places=4)
        self.assertEqual(valores["F17"], "V")


if __name__ == "__main__":
    unittest.main()
